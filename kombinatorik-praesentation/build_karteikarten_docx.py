"""Build per-person Karteikarten (index cards) from the Sprechtext content.

Each output file contains only one presenter's lines, in slide order, one
card per page sized like a physical A6 index card so it can be printed and
cut out directly.

Run with: python3 build_karteikarten_docx.py
"""

import os

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

from build_sprechtext_docx import SLIDES, FULL_NAMES, PERSON_COLOR, DARK

OUT_DIR = os.path.dirname(__file__)
GRAY = RGBColor(0x64, 0x74, 0x8B)

CARD_W = Cm(14.8)  # A6 landscape
CARD_H = Cm(10.5)

FILES = {
    "A": "karteikarten_max.docx",
    "B": "karteikarten_simon.docx",
    "C": "karteikarten_daniel.docx",
}


def build_person_deck(person, out_path):
    cards = []
    for slide in SLIDES:
        texts = [text for p, text in slide["lines"] if p == person]
        if texts:
            cards.append((slide["num"], slide["title"], texts))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.page_width = CARD_W
    section.page_height = CARD_H
    section.orientation = WD_ORIENT.LANDSCAPE
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)

    color = PERSON_COLOR.get(person, DARK)
    full_name = FULL_NAMES[person]

    for i, (num, title, texts) in enumerate(cards):
        head = doc.add_paragraph()
        head.paragraph_format.space_after = Pt(8)
        r = head.add_run(f"Folie {num} – {title}")
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = color

        for text in texts:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(13)
            run.font.color.rgb = DARK

        footer = doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fr = footer.add_run(f"{full_name} · Karte {i + 1}/{len(cards)}")
        fr.italic = True
        fr.font.size = Pt(8)
        fr.font.color.rgb = GRAY

        if i < len(cards) - 1:
            doc.add_page_break()

    doc.save(out_path)
    print(f"Saved {out_path} ({len(cards)} Karten)")


if __name__ == "__main__":
    for person, filename in FILES.items():
        build_person_deck(person, os.path.join(OUT_DIR, filename))
