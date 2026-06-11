"""Reusable building blocks for the Kombinatorik deck."""

import os
from PIL import Image

from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from theme import (
    DEEP_BLUE, MID_BLUE, LIGHT_BLUE, ORANGE, YELLOW, DARK, GRAY, WHITE, BG,
    FONT_HEAD, FONT_BODY,
)

ASSETS_DIR = os.path.join(os.path.dirname(__file__), "assets")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.5)
CONTENT_TOP = Inches(1.35)

PERSON_COLORS = {"A": DEEP_BLUE, "B": ORANGE, "C": MID_BLUE, "Alle": GRAY}


def asset_path(name):
    if not name.endswith(".png"):
        name = f"{name}.png"
    return os.path.join(ASSETS_DIR, name)


def new_slide(prs, bg=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def add_rect(slide, left, top, width, height, fill=None, line=None, line_width=Pt(1),
              shape=MSO_SHAPE.RECTANGLE):
    shp = slide.shapes.add_shape(shape, left, top, width, height)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_width
    return shp


def _set_bullet(paragraph, char="•", color=None):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr", "a:buSzPct"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.set("marL", "228600")
    pPr.set("indent", "-228600")
    if color is not None:
        buClr = pPr.makeelement(qn("a:buClr"), {})
        srgb = pPr.makeelement(qn("a:srgbClr"), {"val": str(color)})
        buClr.append(srgb)
        pPr.append(buClr)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    pPr.append(buFont)
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buChar)


def _no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont", "a:buClr", "a:buSzPct"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    buNone = pPr.makeelement(qn("a:buNone"), {})
    pPr.append(buNone)


def add_title_bar(slide, title, person=None, kicker=None):
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.15), fill=DEEP_BLUE)
    tx = slide.shapes.add_textbox(MARGIN, Inches(0.10), Inches(10.3), Inches(1.0))
    tf = tx.text_frame
    tf.word_wrap = True
    if kicker:
        p0 = tf.paragraphs[0]
        _no_bullet(p0)
        r0 = p0.add_run()
        r0.text = kicker
        r0.font.size = Pt(13)
        r0.font.color.rgb = YELLOW
        r0.font.bold = True
        r0.font.name = FONT_BODY
        p1 = tf.add_paragraph()
    else:
        p1 = tf.paragraphs[0]
    _no_bullet(p1)
    r1 = p1.add_run()
    r1.text = title
    r1.font.size = Pt(30)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.name = FONT_HEAD

    if person:
        add_person_chip(slide, person)
    return tx


def add_person_chip(slide, person):
    color = PERSON_COLORS.get(person, GRAY)
    label = "Alle" if person == "Alle" else f"Person {person}"
    w = Inches(1.55)
    h = Inches(0.42)
    chip = add_rect(slide, SLIDE_W - w - Inches(0.35), Inches(0.36), w, h,
                     fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.adjustments[0] = 0.5
    tf = chip.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _no_bullet(p)
    r = p.add_run()
    r.text = label
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = FONT_BODY


def add_text(slide, text, left, top, width, height, size=Pt(18), color=DARK,
              bold=False, italic=False, align=PP_ALIGN.LEFT, font=FONT_BODY,
              anchor=MSO_ANCHOR.TOP):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _no_bullet(p)
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = size
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = font
    return tx


def add_bullets(slide, items, left, top, width, height, size=Pt(18), color=DARK,
                 bullet_color=ORANGE, space_after=Pt(10), font=FONT_BODY):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = space_after
        _set_bullet(p, color=bullet_color)
        r = p.add_run()
        r.text = item
        r.font.size = size
        r.font.color.rgb = color
        r.font.name = font
    return tx


def add_picture(slide, name, left, top, width=None, height=None):
    return slide.shapes.add_picture(asset_path(name), left, top, width=width, height=height)


def add_picture_fit(slide, name, box_left, box_top, box_width, box_height):
    """Place an image centered within a bounding box, preserving aspect ratio."""
    with Image.open(asset_path(name)) as im:
        iw, ih = im.size
    aspect = iw / ih
    box_aspect = box_width / box_height
    if aspect > box_aspect:
        w = box_width
        h = Emu(int(box_width / aspect))
    else:
        h = box_height
        w = Emu(int(box_height * aspect))
    left = Emu(int(box_left + (box_width - w) / 2))
    top = Emu(int(box_top + (box_height - h) / 2))
    return slide.shapes.add_picture(asset_path(name), left, top, width=w, height=h)


def add_formula(slide, name, left=None, top=None, height=None, width=None,
                 box_left=None, box_top=None, box_width=None, box_height=None):
    """Add a formula image, either at exact position or fit-centered in a box."""
    if box_left is not None:
        return add_picture_fit(slide, name, box_left, box_top, box_width, box_height)
    return slide.shapes.add_picture(asset_path(name), left, top, width=width, height=height)


def add_label_chip(slide, text, left, top, fill=ORANGE, text_color=WHITE, width=Inches(2.4),
                    height=Inches(0.5), size=Pt(16)):
    chip = add_rect(slide, left, top, width, height, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    chip.adjustments[0] = 0.35
    tf = chip.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.1)
    tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _no_bullet(p)
    r = p.add_run()
    r.text = text
    r.font.size = size
    r.font.bold = True
    r.font.color.rgb = text_color
    r.font.name = FONT_BODY
    return chip


def add_panel(slide, left, top, width, height, fill=LIGHT_BLUE, line=MID_BLUE):
    return add_rect(slide, left, top, width, height, fill=fill, line=line, line_width=Pt(1.25),
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE)


def set_notes(slide, text):
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def page_number(slide, n):
    add_text(slide, str(n), SLIDE_W - Inches(0.6), SLIDE_H - Inches(0.45),
             Inches(0.4), Inches(0.35), size=Pt(11), color=GRAY, align=PP_ALIGN.RIGHT)


# Public alias so build_presentation.py can use it via `from slide_helpers import *`
no_bullet = _no_bullet


def add_colored_word(slide, letters_colors, left, top, width, height, size=Pt(40)):
    """Render a word where each letter can have its own color (e.g. ANANAS)."""
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    no_bullet(p)
    for letter, color in letters_colors:
        r = p.add_run()
        r.text = letter
        r.font.size = size
        r.font.bold = True
        r.font.color.rgb = color
        r.font.name = FONT_HEAD
    return tx


def add_circle_badge(slide, text, left, top, diameter, fill, text_color=WHITE, size=Pt(24)):
    circ = add_rect(slide, left, top, diameter, diameter, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = circ.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    no_bullet(p)
    r = p.add_run()
    r.text = text
    r.font.size = size
    r.font.bold = True
    r.font.color.rgb = text_color
    r.font.name = FONT_HEAD
    return circ
